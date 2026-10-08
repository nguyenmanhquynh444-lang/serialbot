"""
excel_writer.py
Ghi kết quả chạy vào một sheet mới trong chính file nguồn, và dựng file báo cáo
riêng khi người dùng chủ động bấm Export.
"""
import logging
import os
import tempfile
import time
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from actions import REASON_LABELS, REASON_NOT_FOUND
from verification import UNVERIFIED, VERIFICATION_LABELS


def atomic_save(workbook, path):
    """A failed ZIP write must leave the existing workbook intact."""
    descriptor, temporary = tempfile.mkstemp(prefix=".serialbot_", suffix=os.path.splitext(path)[1],
                                             dir=os.path.dirname(os.path.abspath(path)))
    os.close(descriptor)
    try:
        workbook.save(temporary)
        with open(temporary, "r+b") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)

STATUS_FILL = {
    "OK": PatternFill("solid", fgColor="1E4620"),
    "ERROR": PatternFill("solid", fgColor="4A1B24"),
    "UNKNOWN": PatternFill("solid", fgColor="4A3E14"),
}
STATUS_FONT = {
    "OK": Font(color="2ECC71"),
    "ERROR": Font(color="FF5C72"),
    "UNKNOWN": Font(color="F1C40F"),
}
HEADER_FILL = PatternFill("solid", fgColor="0F3460")
HEADER_FONT = Font(color="FFFFFF", bold=True)

logger = logging.getLogger(__name__)

RESULT_SHEET_PREFIX = "Ket qua"
# Excel chặn tên sheet dài quá 31 ký tự và không nhận các ký tự này.
MAX_SHEET_NAME = 31
_BAD_SHEET_CHARS = r"[]:*?/\'"


def result_sheet_name(workbook, prefix: str = RESULT_SHEET_PREFIX, now=None) -> str:
    """
    Tên sheet cho phiên chạy này: 'Ket qua 20260812_143052'.

    Có giây trong tên vì chạy lại ngay để thử một dòng là chuyện thường; tới giây
    mà vẫn trùng thì thêm hậu tố, chứ trùng tên là openpyxl tự đổi tên trong im
    lặng và người dùng đi tìm sheet không thấy.
    """
    stamp = (now or datetime.now()).strftime("%Y%m%d_%H%M%S")
    base = "".join(c for c in f"{prefix} {stamp}" if c not in _BAD_SHEET_CHARS)
    base = base[:MAX_SHEET_NAME]

    existing = set(workbook.sheetnames)
    if base not in existing:
        return base
    for n in range(2, 100):
        suffix = f" ({n})"
        candidate = base[:MAX_SHEET_NAME - len(suffix)] + suffix
        if candidate not in existing:
            return candidate
    return base[:MAX_SHEET_NAME]


def find_formula_cells(source_path: str, sheet_name: str | None = None, limit: int = 5) -> list[str]:
    """
    Vài ô đang là công thức trong sheet dữ liệu - rỗng nghĩa là ghi vào file này
    không mất gì.

    Ghi vào file nguồn bằng openpyxl thì GIÁ TRỊ ĐÃ TÍNH của mọi ô công thức bị
    xóa: Excel lưu kèm cả `<f>công thức</f><v>giá trị</v>`, openpyxl chỉ đọc phần
    công thức nên lưu ra là phần giá trị biến mất. Ô đó rỗng cho tới khi ai đó mở
    bằng Excel và lưu lại.

    Nguy ở chỗ excel_reader đọc bằng `data_only=True`: cột serial mà là công thức
    (VLOOKUP sang sheet khác, ghép chuỗi...) thì lần chạy sau bot đọc ra rỗng và
    bỏ qua sạch mọi dòng, trong khi mở file bằng Excel nhìn vẫn thấy đủ số.
    """
    try:
        wb = load_workbook(source_path, read_only=True)
    except Exception:                                       # noqa: BLE001
        return []

    try:
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb[wb.sheetnames[0]]

        found = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    found.append(cell.coordinate)
                    if len(found) >= limit:
                        return found
        return found
    finally:
        wb.close()


class ResultWriter:
    """
    Ghi kết quả vào một sheet mới trong CHÍNH file nguồn - không đẻ file nào thêm.

    Trước đây mỗi phiên chạy tạo một bản copy đầy đủ `<tên>_ket_qua_<giờ>.xlsx`;
    chạy vài lần là thư mục đầy những file gần giống hệt nhau và không còn biết
    file nào là file đang dùng.

    Sheet dữ liệu gốc không bị chạm một ô nào: bảng kết quả tự mang theo số hàng
    Excel và serial nên đối chiếu ngược được. Đây cũng là chỗ sửa một lỗi cũ -
    bản copy ghi vào `wb.active`, tức sheet nào đang mở lúc lưu file, chứ không
    phải sheet người dùng chọn để đọc serial.

    Ghi thẳng vào file nguồn nên phải thử lưu ngay từ đầu (xem `__init__`): file
    đang mở trong Excel thì lưu được lúc cuối phiên là chuyện không có thật, mà
    biết điều đó sau khi đã quét xong vài trăm serial thì đã muộn.
    """

    HEADERS = [
        "Hang Excel", "Serial", "Trang thai", "Phan hoi app",
        "Thoi gian", "Ly do", "Anh chup", "Muc xac nhan",
    ]
    WIDTHS = (12, 24, 14, 44, 12, 22, 40, 24)

    def __init__(
        self,
        source_path: str,
        data_sheet: str | None = None,
        save_every: int = 50,
        result_sheet: str | None = None,
        save_initial: bool = True,
        auto_save: bool = True,
    ):
        self.path = source_path
        # File có macro thì phải nói rõ là giữ: mặc định openpyxl bỏ phần VBA đi,
        # lưu xong là file .xlsm còn nguyên đuôi nhưng sạch macro.
        keep_vba = source_path.lower().endswith((".xlsm", ".xltm"))
        self.wb = load_workbook(source_path, keep_vba=keep_vba)
        self.sheet_name = result_sheet or result_sheet_name(self.wb)
        existing = self.sheet_name in self.wb.sheetnames
        self.ws = self.wb[self.sheet_name] if existing else self.wb.create_sheet(self.sheet_name)
        self.data_sheet = data_sheet
        self.save_every = save_every
        self.auto_save = auto_save
        self._pending = 1
        self._next_row = self.ws.max_row + 1 if existing else 2
        self._row_index = {self.ws.cell(r, 1).value: r for r in range(2, self._next_row)}
        self._last_save = time.monotonic()

        for index, header in enumerate(self.HEADERS, 1):
            self.ws.cell(1, index, header)
        for cell in self.ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for index, width in enumerate(self.WIDTHS, start=1):
            self.ws.column_dimensions[get_column_letter(index)].width = width
        self.ws.freeze_panes = "A2"

        # Lưu ngay khi chưa chạy dòng nào: quyền ghi và khóa file lộ ra ở đây,
        # lúc bên gọi còn hủy phiên được, chứ không phải ở phút thứ mười.
        if save_initial:
            atomic_save(self.wb, self.path)
            self._pending = 0

    # ------------------------------------------------------------------ ghi

    def add_result(
        self,
        row: int,
        serial: str,
        status: str,
        response_text: str,
        screenshot_path: str | None,
        reason: str = "",
        verification: str = UNVERIFIED,
    ):
        values = [
            row,
            serial or "",
            status,
            response_text or "",
            datetime.now().strftime("%H:%M:%S"),
            REASON_LABELS.get(reason, reason or ""),
            screenshot_path or "",
            VERIFICATION_LABELS.get(verification, VERIFICATION_LABELS[UNVERIFIED]),
        ]
        fill = STATUS_FILL.get(status)
        font = STATUS_FONT.get(status)
        target_row = self._row_index.get(row, self._next_row)
        for index, value in enumerate(values, start=1):
            cell = self.ws.cell(row=target_row, column=index, value=value)
            if index == 3:              # chỉ tô cột Trạng thái, tô cả dòng thì rối
                if fill:
                    cell.fill = fill
                if font:
                    cell.font = font

        if row not in self._row_index:
            self._row_index[row] = target_row
            self._next_row += 1
        self._pending += 1
        if self.auto_save and (self._pending >= self.save_every or time.monotonic() - self._last_save >= 30):
            self.save()

    def add_cell(self, row: int, column_letter: str, value):
        """
        Ghi một ô bất kỳ cho bước 'write_cell' của kịch bản.

        Đây là chỗ DUY NHẤT chạm vào sheet dữ liệu gốc, và chỉ chạy khi kịch bản
        khai đích danh - bản chất của bước đó là điền ngược vào bảng của người
        dùng, không có nghĩa gì nếu ghi sang sheet kết quả.
        """
        try:
            col_idx = column_index_from_string(column_letter.strip().upper())
        except Exception:
            return
        self._data_worksheet().cell(row=row, column=col_idx, value=value)
        self._pending += 1

    def _data_worksheet(self):
        """Sheet chứa serial - cùng quy tắc với excel_reader: rỗng = sheet đầu."""
        if self.data_sheet and self.data_sheet in self.wb.sheetnames:
            return self.wb[self.data_sheet]
        return self.wb[self.wb.sheetnames[0]]

    # ---------------------------------------------------------------- lưu

    def save(self):
        if self._pending == 0:
            return
        try:
            atomic_save(self.wb, self.path)
            self._pending = 0
            self._last_save = time.monotonic()
        except PermissionError:
            # Gần như luôn là file đang mở trong Excel. Giữ lại trong bộ nhớ và
            # thử tiếp ở lần lưu sau, đừng làm đổ phiên chạy vì một lần khóa file.
            logger.warning(
                "Chưa lưu được '%s' (file đang mở?), giữ %d dòng chờ lần lưu sau",
                self.path, self._pending,
            )

    def finalize(self):
        atomic_save(self.wb, self.path)
        self._pending = 0
        self._last_save = time.monotonic()


def write_report(path: str, records: list[dict], summary: dict):
    """
    records: [{"time","row","serial","status","message"}, ...] (all rows, not just errors)
    summary: {"total","done","success","error","unknown","started_at","finished_at","elapsed_seconds","rate_per_min"}
    """
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    _write_summary_sheet(ws_summary, summary)

    ws_detail = wb.create_sheet("Chi tiet")
    _write_records_sheet(ws_detail, records)

    ws_errors = wb.create_sheet("Loi")
    error_records = [r for r in records if r["status"] != "OK"]
    _write_records_sheet(ws_errors, error_records)

    # Sheet riêng cho serial hệ thống không có: đây là danh sách phải mang đi
    # tra lại, khác hẳn các lỗi còn lại vốn chỉ cần chạy lại bot.
    ws_not_found = wb.create_sheet("Khong tim thay")
    _write_records_sheet(
        ws_not_found, [r for r in records if r.get("reason") == REASON_NOT_FOUND]
    )

    wb.save(path)


def write_pick_report(
    path: str,
    bills: list[dict],
    records: list[dict],
    summary: dict,
    bills_sheet: str = "Van don",
    bill_header: str = "Ma van don",
    ref_header: str = "Ma van don (MVD)",
):
    """
    Báo cáo cho các chế độ chạy theo phiên (xuất kho, nhập hàng).

    Khác báo cáo thường ở sheet "Ma OR"/"Check-in": ở những luồng này đơn vị
    công việc là cả phiên chứ không phải serial - cần nhìn ngay phiên nào xong,
    phiên nào dở dang và dở tới đâu, vì phiên dở nghĩa là còn một phiên đang mở
    trên máy ảo.

    bills: [{"bill","ref","session_code","expected","picked","status","message"}, ...]
    """
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    _write_summary_sheet(ws_summary, summary)

    ws_bills = wb.create_sheet(bills_sheet)
    _write_bills_sheet(ws_bills, bills, bill_header, ref_header)

    ws_detail = wb.create_sheet("Chi tiet")
    _write_records_sheet(ws_detail, records)

    ws_errors = wb.create_sheet("Loi")
    _write_records_sheet(ws_errors, [r for r in records if r.get("status") != "OK"])

    wb.save(path)


def _write_bills_sheet(
    ws,
    bills: list[dict],
    bill_header: str = "Ma van don",
    ref_header: str = "Ma van don (MVD)",
):
    # Cột MVĐ chỉ có ở luồng xuất kho theo mã OR, và chỉ khi file nguồn khai cột
    # đó - báo cáo của luồng nhập hàng không phải gánh thêm một cột luôn rỗng.
    has_ref = any(bill.get("ref") for bill in bills)

    headers = [bill_header]
    if has_ref:
        headers.append(ref_header)
    headers += ["Ma phien", "Can xuat", "Da xuat", "Trang thai", "Thong bao"]
    status_col = len(headers) - 1

    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for bill in bills:
        expected = bill.get("expected", 0)
        picked = bill.get("picked", 0)
        # Trạng thái tính lại từ số lượng chứ không tin cột status: xuất thiếu mà
        # vẫn ghi OK thì người đọc báo cáo không biết còn phiên đang dở.
        status = bill.get("status", "OK")
        if status == "OK" and picked != expected:
            status = "ERROR"
        row = [bill.get("bill", "")]
        if has_ref:
            row.append(bill.get("ref", ""))
        row += [
            bill.get("session_code", ""),
            expected,
            picked,
            status,
            bill.get("message", ""),
        ]
        ws.append(row)
        font = STATUS_FONT.get(status)
        if font:
            ws.cell(row=ws.max_row, column=status_col).font = font

    widths = [20] + ([20] if has_ref else []) + [26, 12, 12, 14, 50]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"


def _write_summary_sheet(ws, summary: dict):
    rows = [
        ("Tong so serial", summary.get("total", 0)),
        ("Da xu ly", summary.get("done", 0)),
        ("Thanh cong", summary.get("success", 0)),
        ("Da xac nhan", summary.get("confirmed", 0)),
        ("Khong phat hien loi", summary.get("inferred", 0)),
        ("Chua xac minh", summary.get("unverified", 0)),
        ("Loi", summary.get("error", 0)),
        ("  - trong do: khong tim thay serial", summary.get("not_found", 0)),
        ("Khong xac dinh", summary.get("unknown", 0)),
        (
            "Ty le thanh cong",
            f"{summary.get('success', 0) / summary.get('done', 1) * 100:.1f}%"
            if summary.get("done")
            else "0%",
        ),
        ("Bat dau", summary.get("started_at", "")),
        ("Ket thuc", summary.get("finished_at", "")),
        ("Thoi gian chay", summary.get("elapsed_seconds", 0)),
        ("Toc do (serial/phut)", summary.get("rate_per_min", 0)),
    ]
    ws.append(["Chi so", "Gia tri"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for label, value in rows:
        ws.append([label, value])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 24


def _write_records_sheet(ws, records: list[dict]):
    ws.append(["Thoi gian", "Hang Excel", "Serial", "Trang thai", "Ly do", "Thong bao", "Muc xac nhan"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for record in records:
        status = record.get("status", "")
        reason = record.get("reason", "")
        ws.append([
            record.get("time", ""),
            record.get("row", ""),
            record.get("serial", ""),
            status,
            REASON_LABELS.get(reason, reason or ""),
            record.get("message", ""),
            VERIFICATION_LABELS.get(record.get("verification", UNVERIFIED), VERIFICATION_LABELS[UNVERIFIED]),
        ])
        font = STATUS_FONT.get(status)
        if font:
            ws.cell(row=ws.max_row, column=4).font = font

    for col_idx, width in enumerate((12, 12, 24, 12, 22, 50, 24), start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
