"""
Đọc file "mã vận đơn / serial" và gộp thành từng phiên pick.

Cái bẫy lớn nhất là ô gộp: openpyxl chỉ trả giá trị ở dòng trên cùng của vùng
gộp, các dòng dưới là None. Không kéo giá trị xuống thì mỗi vận đơn chỉ nhận
được đúng một serial - mà file kết quả trông vẫn bình thường.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from openpyxl import Workbook

from excel_reader import ExcelReadError, read_bill_groups


def build(tmp_path, rows, merges=(), header=True):
    """rows = [(bill, serial), ...]; bill None nghĩa là ô trống (nằm trong vùng gộp)."""
    wb = Workbook()
    ws = wb.active
    start = 1
    if header:
        ws["A1"], ws["B1"] = "mã vd", "số seri"
        start = 2
    for i, (bill, serial) in enumerate(rows):
        if bill is not None:
            ws.cell(row=start + i, column=1, value=bill)
        if serial is not None:
            ws.cell(row=start + i, column=2, value=serial)
    for m in merges:
        ws.merge_cells(m)
    path = tmp_path / "bills.xlsx"
    wb.save(path)
    return str(path)


def test_preview_rejects_header_even_when_selected_range_includes_it(tmp_path):
    path = build(tmp_path, [("MaOR", "Serial"), ("OR1", "S1")], header=False)
    groups, warnings = read_bill_groups(path, None, "A", "B", 1)
    assert [group["bill"] for group in groups] == ["OR1"]
    assert groups[0]["rows"][0]["value"] == "S1"
    assert warnings


def test_merged_bill_cell_is_carried_down_to_every_serial(tmp_path):
    """Đúng cấu trúc file thật: mã vận đơn là ô gộp trải trên các dòng serial."""
    path = build(
        tmp_path,
        [("8VGR23072026", "BSS-001"), (None, "BSS-002"), (None, "BSS-003"),
         ("24VGR23072026", "BSS-004"), (None, "BSS-005")],
        merges=("A2:A4", "A5:A6"),
    )

    groups, warnings = read_bill_groups(path, None, "A", "B", 2)

    assert [g["bill"] for g in groups] == ["8VGR23072026", "24VGR23072026"]
    assert [len(g["rows"]) for g in groups] == [3, 2]
    assert [r["value"] for r in groups[0]["rows"]] == ["BSS-001", "BSS-002", "BSS-003"]
    assert warnings == []


def test_row_numbers_point_back_at_the_real_excel_rows(tmp_path):
    """Ghi kết quả và báo lỗi đều theo số dòng, lệch một dòng là tra nhầm chỗ."""
    path = build(tmp_path, [("A1", "S1"), (None, "S2"), ("B1", "S3")], merges=("A2:A3",))

    groups, _ = read_bill_groups(path, None, "A", "B", 2)

    assert groups[0]["first_row"] == 2
    assert [r["row"] for r in groups[0]["rows"]] == [2, 3]
    assert groups[1]["first_row"] == 4


def test_serial_with_stray_space_is_cleaned(tmp_path):
    """File thật có 'BSS- 26063041023958' - dấu cách lọt vào giữa."""
    path = build(tmp_path, [("A1", "BSS- 2606304"), (None, "BSS-2606305")], merges=("A2:A3",))

    groups, _ = read_bill_groups(path, None, "A", "B", 2)

    assert [r["value"] for r in groups[0]["rows"]] == ["BSS-2606304", "BSS-2606305"]


def test_serial_before_any_bill_is_reported_not_silently_grouped(tmp_path):
    """Gán bừa vào vận đơn nào đó là xuất kho nhầm đơn."""
    path = build(tmp_path, [(None, "S0"), ("A1", "S1")])

    groups, warnings = read_bill_groups(path, None, "A", "B", 2)

    assert [g["bill"] for g in groups] == ["A1"]
    assert [r["value"] for r in groups[0]["rows"]] == ["S1"]
    assert any("chưa có mã vận đơn" in w for w in warnings)


def test_blank_serial_rows_do_not_break_a_group(tmp_path):
    path = build(tmp_path, [("A1", "S1"), (None, None), (None, "S2")], merges=("A2:A4",))

    groups, _ = read_bill_groups(path, None, "A", "B", 2)

    assert len(groups) == 1
    assert [r["row"] for r in groups[0]["rows"]] == [2, 4]


def test_duplicate_serial_is_warned(tmp_path):
    path = build(tmp_path, [("A1", "S1"), (None, "S1")], merges=("A2:A3",))

    _, warnings = read_bill_groups(path, None, "A", "B", 2)

    assert any("trùng" in w for w in warnings)


def test_same_bill_split_into_two_blocks_is_warned(tmp_path):
    """File bị sắp lẫn: chạy tiếp sẽ mở hai phiên pick cho cùng một vận đơn."""
    path = build(tmp_path, [("A1", "S1"), ("B1", "S2"), ("A1", "S3")])

    groups, warnings = read_bill_groups(path, None, "A", "B", 2)

    assert len(groups) == 3
    assert any("lại tách ra cụm khác" in w for w in warnings)


def test_same_column_for_bill_and_serial_is_rejected(tmp_path):
    path = build(tmp_path, [("A1", "S1")])

    with pytest.raises(ExcelReadError, match="phải khác nhau"):
        read_bill_groups(path, None, "A", "A", 2)


# Dữ liệu thật người dùng gửi: mã lặp xuống từng dòng (không gộp ô), và một cột
# serial thứ hai để soát chéo - cùng tập serial nhưng khác thứ tự.
REAL_BLOCKS = [
    ("195PIVGR27072026",
     ["E52771260430D01999", "E52771260509N00734", "E52771260412D01939",
      "E52771260312N01180", "E52771260313N01474"],
     ["E52771260509N00734", "E52771260312N01180", "E52771260412D01939",
      "E52771260313N01474", "E52771260430D01999"]),
    ("480PIVGR17072026",
     ["E52771260507N02300", "E52771260512N01997", "E52771260528N01733",
      "E52771260615D03222", "E52771260529D01821"],
     ["E52771260529D01821", "E52771260512N01997", "E52771260528N01733",
      "E52771260615D03222", "E52771260507N02300"]),
]


def build_real(tmp_path, cross_check_column=False):
    wb = Workbook()
    ws = wb.active
    ws.append(["Ma van don", "Serial"] + (["Doi chieu"] if cross_check_column else []))
    for bill, serials, other in REAL_BLOCKS:
        for i, serial in enumerate(serials):
            row = [bill, serial]
            if cross_check_column:
                row.append(other[i])
            ws.append(row)
    path = tmp_path / "real.xlsx"
    wb.save(path)
    return str(path)


def test_bill_code_repeated_on_every_row_groups_the_same_as_a_merged_cell(tmp_path):
    """Layout người dùng chốt dùng: mã lặp lại thay vì gộp ô."""
    groups, warnings = read_bill_groups(build_real(tmp_path), None, "A", "B", 2)

    assert [g["bill"] for g in groups] == ["195PIVGR27072026", "480PIVGR17072026"]
    assert [len(g["rows"]) for g in groups] == [5, 5]
    assert groups[0]["rows"][0] == {"row": 2, "value": "E52771260430D01999"}
    assert groups[1]["first_row"] == 7
    assert warnings == []


def test_extra_cross_check_column_is_ignored(tmp_path):
    """Cột soát chéo cùng tập serial khác thứ tự - bot chỉ đọc cột được chỉ định."""
    groups, warnings = read_bill_groups(build_real(tmp_path, True), None, "A", "B", 2)

    assert [len(g["rows"]) for g in groups] == [5, 5]
    assert warnings == []


def test_bill_column_pointing_at_serials_is_warned(tmp_path):
    """
    Cái nhầm im lặng nhất: trỏ 'Cột mã vận đơn' vào một cột serial. File đọc
    trơn tru, hộp xem trước hiện một đống vận đơn trông như thật, tới lúc bot
    tìm phiên pick mới chết.
    """
    path = build_real(tmp_path, cross_check_column=True)

    groups, warnings = read_bill_groups(path, None, "C", "B", 2)

    assert len(groups) == 10                      # mỗi serial thành một "vận đơn"
    assert any("trỏ nhầm" in w for w in warnings)


# --- cột MVĐ đi kèm mã OR ----------------------------------------------------

def build_or_template(tmp_path, rows):
    """rows = [(or_code, mvd, serial), ...] - đúng template trong bản quay màn hình."""
    wb = Workbook()
    ws = wb.active
    ws.append(["OR", "MVĐ", "Seri"])
    for row in rows:
        ws.append(list(row))
    path = tmp_path / "xuat_kho.xlsx"
    wb.save(path)
    return str(path)


def test_mvd_is_carried_alongside_each_or_code(tmp_path):
    path = build_or_template(tmp_path, [
        ("ORVINEKGK4T1560", "581PIPTE03082026", "E52771260715N01954"),
        ("ORVINT8L5UY9235", "580PIPTE03082026", "E52771260715N01854"),
    ])

    groups, warnings = read_bill_groups(path, None, "A", "C", 2, ref_column="B")

    assert [g["bill"] for g in groups] == ["ORVINEKGK4T1560", "ORVINT8L5UY9235"]
    assert [g["ref"] for g in groups] == ["581PIPTE03082026", "580PIPTE03082026"]
    assert groups[0]["rows"][0]["ref"] == "581PIPTE03082026"
    assert warnings == []


def test_mvd_column_pointing_at_the_serial_column_is_rejected(tmp_path):
    path = build_or_template(tmp_path, [("OR1", "MVD1", "S1")])

    with pytest.raises(ExcelReadError, match="tham chiếu"):
        read_bill_groups(path, None, "A", "C", 2, ref_column="C")


def test_one_or_code_with_two_different_mvd_is_warned(tmp_path):
    """Báo cáo chỉ giữ được một MVĐ cho mỗi phiên - nói trước còn hơn ghi thiếu."""
    path = build_or_template(tmp_path, [
        ("OR1", "MVD-A", "S1"),
        ("OR1", "MVD-B", "S2"),
    ])

    groups, warnings = read_bill_groups(path, None, "A", "C", 2, ref_column="B")

    assert len(groups) == 1
    assert any("nhiều mã vận đơn" in w for w in warnings)


def test_leaving_the_mvd_column_empty_still_reads_the_file(tmp_path):
    """File hai cột đời cũ: không khai cột tham chiếu thì không có khóa 'ref' thừa."""
    path = build_or_template(tmp_path, [("OR1", "MVD1", "S1")])

    groups, _ = read_bill_groups(path, None, "A", "C", 2)

    assert groups[0]["ref"] == ""
    assert "ref" not in groups[0]["rows"][0]


def test_warnings_use_the_name_of_the_code_being_run(tmp_path):
    """Gọi 'mã vận đơn' trong lúc người dùng đang chạy theo mã OR là chỉ sai chỗ."""
    path = build_or_template(tmp_path, [(None, None, "S0"), ("OR1", "MVD1", "S1")])

    _, warnings = read_bill_groups(path, None, "A", "C", 2, code_label="mã OR")

    assert any("chưa có mã OR nào" in w for w in warnings)


def test_end_row_limits_how_far_it_reads(tmp_path):
    path = build(tmp_path, [("A1", "S1"), (None, "S2"), (None, "S3")], merges=("A2:A4",))

    groups, _ = read_bill_groups(path, None, "A", "B", 2, end_row=3)

    assert [r["value"] for r in groups[0]["rows"]] == ["S1", "S2"]
