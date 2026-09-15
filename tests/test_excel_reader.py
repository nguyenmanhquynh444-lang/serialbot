import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from openpyxl import Workbook

from excel_reader import (
    clean_serial,
    guess_bill_column,
    guess_serial_column,
    read_serials,
)


def test_clean_serial_strips_whitespace_and_control_chars():
    assert clean_serial(" AB\t12\n34 ") == "AB1234"


def test_clean_serial_handles_float_like_excel_numbers():
    assert clean_serial(12345.0) == "12345"


def test_clean_serial_handles_vietnamese_diacritics():
    assert clean_serial("Đ123đ") == "D123d"


def test_clean_serial_none_and_empty():
    assert clean_serial(None) == ""
    assert clean_serial("") == ""


def test_clean_serial_trailing_dot_zero_string():
    assert clean_serial("98765.0") == "98765"


@pytest.fixture
def sample_xlsx(tmp_path):
    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["A", "B", "C", "D", "E", "F", "G", "Serial"])
    for i in range(2, 8):
        ws.cell(row=i, column=8, value=f"SER{i:04d}")
    wb.save(path)
    return str(path)


def test_read_serials_basic(sample_xlsx):
    results = read_serials(sample_xlsx, None, "H", 2)
    assert len(results) == 6
    assert results[0] == {"row": 2, "value": "SER0002"}
    assert results[-1] == {"row": 7, "value": "SER0007"}


def test_read_serials_respects_end_row(sample_xlsx):
    results = read_serials(sample_xlsx, None, "H", 2, end_row=4)
    assert [r["row"] for r in results] == [2, 3, 4]


def test_guess_serial_column_finds_header(sample_xlsx):
    assert guess_serial_column(sample_xlsx, None) == "H"


@pytest.mark.parametrize(
    "header",
    ["Ma van don", "Mã vận đơn", "MÃ VẬN ĐƠN", "Bill", "Mã bill", "Đơn hàng"],
)
def test_guess_bill_column_reads_the_usual_header_spellings(tmp_path, header):
    """Người dùng gõ tiêu đề có dấu, không dấu, hoa thường lẫn lộn."""
    path = tmp_path / "bills.xlsx"
    wb = Workbook()
    wb.active.append([header, "Serial"])
    wb.save(path)

    assert guess_bill_column(str(path), None) == "A"


def test_guess_bill_column_ignores_a_serial_header(sample_xlsx):
    """File chỉ có cột serial thì đừng đoán bừa ra cột mã vận đơn."""
    assert guess_bill_column(sample_xlsx, None) is None


def test_read_serials_missing_file_raises():
    from excel_reader import ExcelReadError
    with pytest.raises(ExcelReadError):
        read_serials("does_not_exist.xlsx", None, "H", 2)
