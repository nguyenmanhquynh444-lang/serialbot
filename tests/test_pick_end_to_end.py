"""
Chạy suốt chuỗi: file Excel ô gộp -> gom nhóm -> PickWorker -> kết quả + báo cáo.

Từng mảnh đã có test riêng, nhưng chỗ hay hỏng nhất là khớp nối giữa chúng:
read_bill_groups trả về dict dạng này, PickWorker mong dạng kia.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pytest

pytest.importorskip("PyQt6")

from openpyxl import Workbook, load_workbook

from excel_reader import read_bill_groups
from excel_writer import write_pick_report
from fake_pick_device import FakePickDevice, Session
from pick_runner import PickConfig
from pick_worker import PickWorker

# Đúng cấu trúc file người dùng gửi: mã vận đơn là ô gộp trải trên các dòng serial.
DATA = [
    ("8VGR23072026", ["BSS-26063041022726", "BSS-26063041023696", "BSS- 26063041023958"]),
    ("24VGR23072026", ["BSS-26063041022744", "BSS-26063041023364"]),
]

# Template xuất kho trong bản quay màn hình: OR / MVĐ / Seri, mỗi dòng một OR.
OR_DATA = [
    ("ORVINEKGK4T1560", "581PIPTE03082026", "E52771260715N01954"),
    ("ORVINT8L5UY9235", "580PIPTE03082026", "E52771260715N01854"),
]


@pytest.fixture
def excel(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "mã vd", "số seri"
    row = 2
    for bill, serials in DATA:
        ws.cell(row=row, column=1, value=bill)
        for i, serial in enumerate(serials):
            ws.cell(row=row + i, column=2, value=serial)
        ws.merge_cells(
            start_row=row, start_column=1, end_row=row + len(serials) - 1, end_column=1
        )
        row += len(serials)
    path = tmp_path / "van_don.xlsx"
    wb.save(path)
    return str(path)


@pytest.fixture
def or_excel(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["OR", "MVĐ", "Seri"])
    for row in OR_DATA:
        ws.append(list(row))
    path = tmp_path / "xuat_kho.xlsx"
    wb.save(path)
    return str(path)


def run_all(groups, device):
    worker = PickWorker(
        groups, "LDPlayer",
        PickConfig(step_timeout=2.0, capture_screenshot=False),
    )
    worker.adb_client = device
    events = {"results": [], "bills": [], "stopped": []}
    worker.result_ready.connect(events["results"].append)
    worker.bill_done.connect(events["bills"].append)
    worker.stopped.connect(lambda m, p: events["stopped"].append(m))
    worker.run()
    return worker, events


def test_excel_to_finished_run(qapp, excel, tmp_path):
    groups, warnings = read_bill_groups(excel, None, "A", "B", 2)
    assert warnings == []
    assert [len(g["rows"]) for g in groups] == [3, 2]

    device = FakePickDevice([
        Session("8VGR23072026", "PL_001", 3),
        Session("24VGR23072026", "PL_002", 2),
    ])
    worker, events = run_all(groups, device)

    assert events["stopped"] == []
    assert worker.done_bills == ["8VGR23072026", "24VGR23072026"]

    # Dấu cách lọt giữa serial trong file gốc phải được dọn trước khi gửi đi.
    sent = [text for state, text in device.typed if state == "serial"]
    assert "BSS-26063041023958" in sent
    assert not any(" " in s for s in sent)

    # Số dòng Excel đi được tới tận kết quả, để ghi ngược đúng chỗ.
    assert [r["row"] for r in events["results"]] == [2, 3, 4, 5, 6]

    report = str(tmp_path / "bao_cao.xlsx")
    write_pick_report(report, events["bills"], events["results"], {"total": 5, "done": 5})
    rows = list(load_workbook(report)["Van don"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["8VGR23072026", "24VGR23072026"]
    assert [r[4] for r in rows] == ["OK", "OK"]


def test_excel_count_mismatch_stops_before_the_second_bill(qapp, excel):
    """App báo vận đơn 2 cần 9 món, Excel chỉ có 2 - phải dừng, đơn 1 vẫn giữ."""
    groups, _ = read_bill_groups(excel, None, "A", "B", 2)
    device = FakePickDevice([
        Session("8VGR23072026", "PL_001", 3),
        Session("24VGR23072026", "PL_002", 9),
    ])

    worker, events = run_all(groups, device)

    assert worker.done_bills == ["8VGR23072026"]
    assert "app báo cần 9" in events["stopped"][0]
    assert [b["bill"] for b in events["bills"]] == ["8VGR23072026"]


def test_three_column_template_runs_and_keeps_the_mvd(qapp, or_excel, tmp_path):
    """
    Template trong bản quay màn hình chạy trọn: cột A là mã bot gõ, cột B chỉ đi
    kèm. Lẫn hai cột là bot tra một mã app không nhận.
    """
    groups, warnings = read_bill_groups(
        or_excel, None, "A", "C", 2, ref_column="B", code_label="mã OR",
    )

    assert warnings == []
    assert [g["bill"] for g in groups] == ["ORVINEKGK4T1560", "ORVINT8L5UY9235"]
    assert [g["ref"] for g in groups] == ["581PIPTE03082026", "580PIPTE03082026"]

    device = FakePickDevice([
        Session("ORVINEKGK4T1560", "PL260804PTD2T_PTD_581", 1),
        Session("ORVINT8L5UY9235", "PL260804PTD2T_PTD_580", 1),
    ])
    worker, events = run_all(groups, device)

    assert events["stopped"] == []
    assert worker.done_bills == ["ORVINEKGK4T1560", "ORVINT8L5UY9235"]

    # Chỉ mã OR và serial được gửi vào app; MVĐ không đi vào ô nhập nào.
    sent = [text for _, text in device.typed]
    assert "581PIPTE03082026" not in sent
    assert sent.count("T001") == 1              # tote chỉ đăng ký một lần
    assert [t for s, t in device.typed if s == "list"] == [
        "ORVINEKGK4T1560", "ORVINT8L5UY9235"
    ]

    report = str(tmp_path / "bao_cao.xlsx")
    write_pick_report(
        report, events["bills"], events["results"], {"total": 2, "done": 2},
        bills_sheet="Ma OR", bill_header="Ma OR",
    )
    rows = list(load_workbook(report)["Ma OR"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["ORVINEKGK4T1560", "ORVINT8L5UY9235"]
    assert [r[1] for r in rows] == ["581PIPTE03082026", "580PIPTE03082026"]
