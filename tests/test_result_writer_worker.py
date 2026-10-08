import threading

from openpyxl import Workbook, load_workbook

import excel_writer
from result_writer_worker import ResultWriterWorker
from result_writer_worker import ReportWriterWorker
import result_writer_worker


def source(tmp_path):
    path = str(tmp_path / "source.xlsx")
    wb = Workbook()
    wb.active.append(["Serial"])
    wb.active.append(["S1"])
    wb.save(path)
    return path


def record(status="OK"):
    return {"row": 2, "serial": "S1", "status": status, "verification": "confirmed", "response_text": "accepted"}


def test_workbook_is_saved_off_gui_thread_and_replayed_without_duplicate_rows(tmp_path, monkeypatch, qapp):
    path = source(tmp_path)
    gui_thread = threading.get_ident()
    saved_in = []
    real_save = excel_writer.atomic_save
    def tracked_save(wb, target):
        saved_in.append(threading.get_ident())
        real_save(wb, target)
    monkeypatch.setattr(excel_writer, "atomic_save", tracked_save)
    for status in ("OK", "ERROR"):
        worker = ResultWriterWorker(path, result_sheet="Ket qua test")
        worker.start()
        worker.add_record(record(status))
        worker.finalize()
        assert worker.wait(5000)
        assert worker.saved
    wb = load_workbook(path)
    assert wb["Ket qua test"].max_row == 2
    assert wb["Ket qua test"].cell(2, 3).value == "ERROR"
    assert all(thread != gui_thread for thread in saved_in)
    wb.close()


def test_atomic_replace_failure_leaves_source_intact_and_can_retry(tmp_path, monkeypatch, qapp):
    path = source(tmp_path)
    before = open(path, "rb").read()
    replace = excel_writer.os.replace
    monkeypatch.setattr(excel_writer.os, "replace", lambda *args: (_ for _ in ()).throw(PermissionError("locked")))
    worker = ResultWriterWorker(path, result_sheet="Ket qua test")
    worker.start()
    worker.add_record(record())
    worker.finalize()
    assert worker.wait(5000)
    assert not worker.saved
    assert open(path, "rb").read() == before
    monkeypatch.setattr(excel_writer.os, "replace", replace)
    retry = ResultWriterWorker(path, result_sheet="Ket qua test")
    retry.start()
    retry.add_record(record())
    retry.finalize()
    assert retry.wait(5000)
    assert retry.saved
    assert not list(tmp_path.glob(".serialbot_*"))


def test_report_export_owns_a_snapshot_and_runs_off_gui_thread(tmp_path, monkeypatch, qapp):
    path = str(tmp_path / "report.xlsx")
    gui_thread = threading.get_ident()
    threads = []
    original = result_writer_worker.write_report
    def tracked(*args, **kwargs):
        threads.append(threading.get_ident())
        return original(*args, **kwargs)
    monkeypatch.setattr(result_writer_worker, "write_report", tracked)
    records = [record()]
    worker = ReportWriterWorker(path, records, {"total": 1, "done": 1})
    records[0]["status"] = "ERROR"
    worker.start()
    assert worker.wait(5000)
    assert worker.error is None
    assert threads[0] != gui_thread
    workbook = load_workbook(path)
    assert workbook["Chi tiet"].cell(2, 4).value == "OK"
    workbook.close()
