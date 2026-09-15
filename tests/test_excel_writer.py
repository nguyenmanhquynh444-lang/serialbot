import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from openpyxl import Workbook, load_workbook

from excel_writer import (
    ResultWriter,
    find_formula_cells,
    result_sheet_name,
    write_report,
)


def make_source(tmp_path, sheets=("Sheet",)):
    path = tmp_path / "source.xlsx"
    wb = Workbook()
    wb.active.title = sheets[0]
    for name in sheets[1:]:
        wb.create_sheet(name)
    ws = wb[sheets[0]]
    ws.append(["A", "B", "C", "D", "E", "F", "G", "Serial"])
    ws.append(["", "", "", "", "", "", "", "SER0001"])
    wb.save(path)
    return str(path)


# --- không đẻ file nào thêm --------------------------------------------------

def test_writes_into_the_source_file_and_creates_nothing_else(tmp_path):
    source = make_source(tmp_path)
    before = set(os.listdir(tmp_path))

    writer = ResultWriter(source)
    writer.add_result(2, "SER0001", "OK", "matched", None)
    writer.finalize()

    assert set(os.listdir(tmp_path)) == before


def test_results_land_in_a_new_sheet(tmp_path):
    source = make_source(tmp_path)
    writer = ResultWriter(source)
    writer.add_result(2, "SER0001", "ERROR", "invalid", "shot.png", "NOT_FOUND")
    writer.finalize()

    wb = load_workbook(source)
    assert writer.sheet_name in wb.sheetnames
    ws = wb[writer.sheet_name]
    assert [c.value for c in ws[1]] == ResultWriter.HEADERS
    assert ws.cell(row=2, column=1).value == 2              # hàng Excel
    assert ws.cell(row=2, column=2).value == "SER0001"
    assert ws.cell(row=2, column=3).value == "ERROR"
    assert ws.cell(row=2, column=4).value == "invalid"
    assert ws.cell(row=2, column=6).value == "Không tìm thấy serial"
    assert ws.cell(row=2, column=7).value == "shot.png"
    wb.close()


def test_source_sheet_is_left_untouched(tmp_path):
    """Bot chấm sai một dòng cũng không được làm hỏng bảng dữ liệu gốc."""
    source = make_source(tmp_path)
    writer = ResultWriter(source)
    writer.add_result(2, "SER0001", "OK", "matched", None)
    writer.finalize()

    ws = load_workbook(source)["Sheet"]
    assert ws.cell(row=1, column=8).value == "Serial"
    assert ws.cell(row=2, column=8).value == "SER0001"
    assert ws.max_column == 8            # không có cột kết quả nào chèn thêm
    assert ws.cell(row=2, column=9).value is None


def test_other_sheets_survive(tmp_path):
    source = make_source(tmp_path, sheets=("Data", "Ghi chu"))
    writer = ResultWriter(source, data_sheet="Data")
    writer.add_result(2, "SER0001", "OK", "", None)
    writer.finalize()

    names = load_workbook(source).sheetnames
    assert names[:2] == ["Data", "Ghi chu"]
    assert writer.sheet_name in names


# --- mỗi phiên một sheet -----------------------------------------------------

def test_a_second_run_adds_another_sheet_instead_of_overwriting(tmp_path):
    source = make_source(tmp_path)

    first = ResultWriter(source)
    first.add_result(2, "SER0001", "OK", "lan 1", None)
    first.finalize()

    second = ResultWriter(source)
    second.add_result(2, "SER0001", "ERROR", "lan 2", None)
    second.finalize()

    wb = load_workbook(source)
    assert first.sheet_name != second.sheet_name
    assert wb[first.sheet_name].cell(row=2, column=4).value == "lan 1"
    assert wb[second.sheet_name].cell(row=2, column=4).value == "lan 2"
    wb.close()


def test_sheet_name_is_unique_even_within_the_same_second():
    """Chạy lại ngay để thử một dòng là chuyện thường; trùng tên thì openpyxl
    tự đổi trong im lặng và người dùng đi tìm sheet không thấy."""
    wb = Workbook()
    fixed = datetime(2026, 8, 12, 14, 30, 52)

    first = result_sheet_name(wb, now=fixed)
    wb.create_sheet(first)
    second = result_sheet_name(wb, now=fixed)

    assert first == "Ket qua 20260812_143052"
    assert second != first
    assert len(second) <= 31


def test_sheet_name_fits_excel_limits():
    name = result_sheet_name(Workbook(), prefix="Ket qua cua phien chay rat dai")
    assert len(name) <= 31
    assert not set(name) & set(r"[]:*?/\'")


# --- ghi ngược vào bảng gốc chỉ khi kịch bản yêu cầu -------------------------

def test_write_cell_targets_the_data_sheet_not_the_result_sheet(tmp_path):
    source = make_source(tmp_path, sheets=("Data", "Khac"))
    writer = ResultWriter(source, data_sheet="Data")
    writer.add_cell(2, "C", "da xu ly")
    writer.finalize()

    wb = load_workbook(source)
    assert wb["Data"].cell(row=2, column=3).value == "da xu ly"
    assert wb[writer.sheet_name].max_row == 1        # mới chỉ có dòng tiêu đề
    wb.close()


def test_write_cell_falls_back_to_the_first_sheet(tmp_path):
    """Cùng quy tắc với excel_reader: không khai sheet thì lấy sheet đầu."""
    source = make_source(tmp_path, sheets=("Data", "Khac"))
    writer = ResultWriter(source)
    writer.add_cell(2, "C", "x")
    writer.finalize()

    assert load_workbook(source)["Data"].cell(row=2, column=3).value == "x"


# --- hỏng quyền ghi phải lộ ra ngay -----------------------------------------

def test_locked_file_fails_at_construction_not_at_the_end(tmp_path, monkeypatch):
    """
    File đang mở trong Excel thì lưu được lúc cuối phiên là chuyện không có thật.
    Biết sau khi đã quét vài trăm serial thì kết quả mất trắng, nên phải thử lưu
    ngay lúc bên gọi còn hủy phiên được.
    """
    source = make_source(tmp_path)

    import openpyxl.workbook.workbook as wb_module

    def refuse(self, filename):
        raise PermissionError("file dang mo trong Excel")

    monkeypatch.setattr(wb_module.Workbook, "save", refuse)

    with pytest.raises(PermissionError):
        ResultWriter(source)


def test_a_lock_appearing_mid_run_does_not_lose_rows(tmp_path, monkeypatch):
    """Người dùng mở file giữa chừng: giữ lại trong bộ nhớ, lưu ở lần sau."""
    source = make_source(tmp_path)
    writer = ResultWriter(source, save_every=1)

    import openpyxl.workbook.workbook as wb_module

    real_save = wb_module.Workbook.save
    monkeypatch.setattr(
        wb_module.Workbook, "save",
        lambda self, filename: (_ for _ in ()).throw(PermissionError("dang mo")),
    )
    writer.add_result(2, "SER0001", "OK", "trong luc bi khoa", None)
    assert writer._pending == 1                      # chưa mất, đang chờ

    monkeypatch.setattr(wb_module.Workbook, "save", real_save)
    writer.finalize()

    ws = load_workbook(source)[writer.sheet_name]
    assert ws.cell(row=2, column=4).value == "trong luc bi khoa"


# --- những thứ ghi vào file nguồn có thể làm hỏng ---------------------------

def test_formula_cells_are_reported_before_writing(tmp_path):
    """
    openpyxl không giữ được giá trị đã tính của ô công thức, mà excel_reader đọc
    bằng data_only=True. Cột serial là công thức thì lần chạy sau đọc ra rỗng.
    """
    path = tmp_path / "co cong thuc.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Ma", "Serial"])
    ws.append([5, "=VLOOKUP(A2,Tra!A:B,2,0)"])
    wb.save(path)

    assert find_formula_cells(str(path), "Data") == ["B2"]


def test_no_formulas_means_nothing_to_warn_about(tmp_path):
    assert find_formula_cells(make_source(tmp_path), "Sheet") == []


def test_formula_scan_stops_at_the_limit(tmp_path):
    path = tmp_path / "nhieu cong thuc.xlsx"
    wb = Workbook()
    ws = wb.active
    for i in range(20):
        ws.append([f"=ROW()+{i}"])
    wb.save(path)

    assert len(find_formula_cells(str(path), limit=3)) == 3


def test_formula_scan_never_raises_on_a_broken_file(tmp_path):
    """Chỗ này chỉ để cảnh báo - hỏng thì không được chặn cả phiên chạy."""
    broken = tmp_path / "hong.xlsx"
    broken.write_bytes(b"khong phai xlsx")

    assert find_formula_cells(str(broken)) == []


def test_macro_workbook_keeps_its_vba(tmp_path):
    """Mặc định openpyxl vứt phần VBA đi - file .xlsm lưu xong là sạch macro."""
    import zipfile

    plain = tmp_path / "tmp.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(plain)

    xlsm = tmp_path / "co macro.xlsm"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(xlsm, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("xl/vbaProject.bin", b"MACRO-BINARY-PLACEHOLDER")

    writer = ResultWriter(str(xlsm))
    writer.add_result(2, "SER0001", "OK", "", None)
    writer.finalize()

    with zipfile.ZipFile(xlsm) as z:
        assert "xl/vbaProject.bin" in z.namelist()


# --- báo cáo Export (người dùng chủ động, vẫn là file riêng) -----------------

def test_write_report_creates_all_sheets(tmp_path):
    out_path = str(tmp_path / "report.xlsx")
    records = [
        {"time": "10:00:00", "row": 2, "serial": "A1", "status": "OK", "message": "ok"},
        {"time": "10:00:01", "row": 3, "serial": "A2", "status": "ERROR", "message": "bad"},
        {"time": "10:00:02", "row": 4, "serial": "A3", "status": "ERROR",
         "message": "khong tim thay", "reason": "NOT_FOUND"},
    ]
    summary = {"total": 3, "done": 3, "success": 1, "error": 2, "not_found": 1, "unknown": 0}
    write_report(out_path, records, summary)

    wb = load_workbook(out_path)
    assert set(wb.sheetnames) == {"Summary", "Chi tiet", "Loi", "Khong tim thay"}
    assert wb["Loi"].max_row == 3  # header + 2 error rows
    # Sheet riêng chỉ chứa dòng không tìm thấy, để mang đi tra lại.
    assert wb["Khong tim thay"].max_row == 2
    assert wb["Khong tim thay"].cell(row=2, column=3).value == "A3"
    assert wb["Khong tim thay"].cell(row=2, column=5).value == "Không tìm thấy serial"
    wb.close()
