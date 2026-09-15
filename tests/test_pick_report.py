"""
Báo cáo và checkpoint của chế độ xuất kho.

Điểm dễ sai nhất: một vận đơn xuất thiếu mà báo cáo vẫn ghi OK. Đơn thiếu nghĩa
là còn một phiên đang mở trên máy ảo - người đọc báo cáo phải thấy ngay.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook

import session_manager
from excel_writer import write_pick_report

BILLS = [
    {"bill": "VD1", "session_code": "PL_001", "expected": 3, "picked": 3,
     "status": "OK", "message": "Đã xuất 3/3 serial."},
    {"bill": "VD2", "session_code": "PL_002", "expected": 5, "picked": 2,
     "status": "OK", "message": "dừng giữa chừng"},
]
RECORDS = [
    {"time": "10:00:00", "row": 2, "serial": "S1", "status": "OK", "message": "Đã xuất"},
    {"time": "10:00:05", "row": 3, "serial": "S2", "status": "ERROR", "message": "bị từ chối"},
]
SUMMARY = {"total": 8, "done": 5, "success": 5, "error": 0, "unknown": 0}


def sheet_rows(ws):
    return [[c for c in row] for row in ws.iter_rows(min_row=2, values_only=True)]


def test_report_has_a_sheet_per_view(tmp_path):
    path = str(tmp_path / "bao_cao.xlsx")
    write_pick_report(path, BILLS, RECORDS, SUMMARY)

    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Van don", "Chi tiet", "Loi"]


def test_short_picked_bill_is_marked_error_even_if_status_said_ok(tmp_path):
    """Xuất 2/5 mà ghi OK thì người đọc không biết còn phiên đang dở."""
    path = str(tmp_path / "bao_cao.xlsx")
    write_pick_report(path, BILLS, RECORDS, SUMMARY)

    rows = sheet_rows(load_workbook(path)["Van don"])

    assert rows[0][:5] == ["VD1", "PL_001", 3, 3, "OK"]
    assert rows[1][:5] == ["VD2", "PL_002", 5, 2, "ERROR"]


def test_error_sheet_only_keeps_failures(tmp_path):
    path = str(tmp_path / "bao_cao.xlsx")
    write_pick_report(path, BILLS, RECORDS, SUMMARY)

    rows = sheet_rows(load_workbook(path)["Loi"])

    assert [r[2] for r in rows] == ["S2"]


# --- cột mã vận đơn đi kèm mã OR --------------------------------------------

OR_BILLS = [
    {"bill": "ORVINEKGK4T1560", "ref": "581PIPTE03082026", "session_code": "PL_581",
     "expected": 1, "picked": 1, "status": "OK", "message": "Đã xuất 1/1 serial."},
]


def test_mvd_column_appears_next_to_the_or_code(tmp_path):
    """Báo cáo phải tra ngược được về đơn hàng, không chỉ về mã phiên của app."""
    path = str(tmp_path / "bao_cao.xlsx")
    write_pick_report(
        path, OR_BILLS, RECORDS, SUMMARY,
        bills_sheet="Ma OR", bill_header="Ma OR",
    )

    ws = load_workbook(path)["Ma OR"]

    assert [c.value for c in ws[1]][:3] == ["Ma OR", "Ma van don (MVD)", "Ma phien"]
    assert sheet_rows(ws)[0][:4] == ["ORVINEKGK4T1560", "581PIPTE03082026", "PL_581", 1]


def test_no_mvd_column_when_the_file_had_none(tmp_path):
    """Luồng nhập hàng không có MVĐ - đừng bắt nó gánh một cột luôn rỗng."""
    path = str(tmp_path / "bao_cao.xlsx")
    write_pick_report(path, BILLS, RECORDS, SUMMARY)

    headers = [c.value for c in load_workbook(path)["Van don"][1]]

    assert "Ma van don (MVD)" not in headers


# --- checkpoint -------------------------------------------------------------

def test_pick_checkpoint_remembers_finished_bills(tmp_path, monkeypatch):
    monkeypatch.setattr(session_manager, "SESSION_DIR", str(tmp_path))
    monkeypatch.setattr(session_manager, "SESSION_PATH", str(tmp_path / "session.json"))
    excel = str(tmp_path / "nguon.xlsx")

    session_manager.save_session(
        excel, done_rows=[2, 3, 4], done=3, success=3, error=0, unknown=0,
        done_bills=["VD1", "VD2"], mode="pick",
    )
    loaded = session_manager.load_session(excel)

    assert loaded["mode"] == "pick"
    assert loaded["done_bills"] == ["VD1", "VD2"]


def test_old_checkpoints_still_load(tmp_path, monkeypatch):
    """File checkpoint đời trước không có hai khóa mới - không được vỡ."""
    monkeypatch.setattr(session_manager, "SESSION_DIR", str(tmp_path))
    monkeypatch.setattr(session_manager, "SESSION_PATH", str(tmp_path / "session.json"))
    excel = str(tmp_path / "nguon.xlsx")

    session_manager.save_session(excel, [2], 1, 1, 0, 0)
    loaded = session_manager.load_session(excel)

    assert loaded["mode"] == "serial"
    assert loaded["done_bills"] == []
